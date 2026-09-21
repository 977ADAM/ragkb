"""Загрузчики файлов: порядок блоков и то, что из документа нельзя потерять.

Проверяется прежде всего DOCX: в справочниках почти весь смысл лежит
таблицами, а от порядка блоков зависит, к какому разделу они попадут и какую
ссылку на источник увидит пользователь.
"""
from __future__ import annotations

import docx
import pytest

from ragkb.core import loaders
from ragkb.core.documents import section_documents


def _make_docx(path, items) -> None:
    """Собирает docx из последовательности («heading», уровень, текст) и таблиц."""
    document = docx.Document()
    for item in items:
        kind = item[0]
        if kind == "heading":
            document.add_heading(item[2], level=item[1])
        elif kind == "para":
            document.add_paragraph(item[1])
        elif kind == "table":
            rows = item[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for cell_index, value in enumerate(row):
                    table.cell(row_index, cell_index).text = value
    document.save(str(path))


def _sections(path) -> list[tuple[str, str]]:
    """Пары «раздел, текст блока» — так видно, куда попала каждая таблица."""
    loaded = loaders.load(path)
    return [
        (str(document.metadata.get("section") or ""), document.page_content)
        for document in section_documents(loaded)
    ]


def test_table_stays_inside_its_own_section(tmp_path):
    """Таблица принадлежит своему разделу, а не последнему в документе.

    Раньше все таблицы дописывались в конец списка блоков, поэтому таблица
    первого раздела попадала в раздел последний — и ссылка на источник врала.
    """
    path = tmp_path / "spravochnik.docx"
    _make_docx(
        path,
        [
            ("heading", 1, "Первый раздел"),
            ("table", [["Продукт", "Задача"], ["AdSmart", "Охват"]]),
            ("heading", 1, "Второй раздел"),
            ("table", [["Продукт", "Задача"], ["In-Stream", "Видео"]]),
        ],
    )

    sections = _sections(path)
    by_section = dict(sections)

    assert "AdSmart" in by_section["Первый раздел"]
    assert "In-Stream" in by_section["Второй раздел"]
    # Главная проверка: содержимое первой таблицы не уехало во второй раздел.
    assert "AdSmart" not in by_section["Второй раздел"]


def test_single_row_table_is_not_lost(tmp_path):
    """Однострочная таблица — это абзац в рамке, а не таблица без данных.

    Раньше её первая строка считалась шапкой, а строк с данными не было,
    поэтому текст пропадал целиком (в разобранном справочнике так терялось
    вступление документа).
    """
    path = tmp_path / "vstuplenie.docx"
    _make_docx(
        path,
        [
            ("heading", 1, "Назначение"),
            ("table", [["Внутренний мастер-документ для отдела продаж."]]),
        ],
    )

    text = "\n".join(block.text for block in loaders.load(path).blocks)

    assert "мастер-документ для отдела продаж" in text


def test_multi_row_table_becomes_column_pairs(tmp_path):
    """Строка таблицы читается как «колонка: значение» — так её ищет BM25."""
    path = tmp_path / "tablica.docx"
    _make_docx(
        path,
        [("table", [["Продукт", "Задача"], ["AdSmart", "Охват"]])],
    )

    blocks = loaders.load(path).blocks

    assert [block.text for block in blocks] == ["Продукт: AdSmart; Задача: Охват"]
    assert blocks[0].kind == "table"


def test_heading_styles_become_headings(tmp_path):
    """Заголовки задают разделы: по ним строится путь в ссылке на источник."""
    path = tmp_path / "zagolovki.docx"
    _make_docx(
        path,
        [
            ("heading", 1, "Продукты"),
            ("heading", 2, "Баннеры"),
            ("para", "Текст про баннеры."),
        ],
    )

    blocks = loaders.load(path).blocks
    headings = [(block.level, block.text) for block in blocks if block.kind == "heading"]

    assert headings == [(1, "Продукты"), (2, "Баннеры")]
    assert [document.metadata["section"] for document in section_documents(loaders.load(path))] == [
        "Продукты > Баннеры"
    ]


def test_unsupported_extension_is_reported(tmp_path):
    path = tmp_path / "dannye.xyz"
    path.write_text("что-то", encoding="utf-8")

    with pytest.raises(loaders.UnsupportedFormat) as exc:
        loaders.load(path)

    assert ".xyz" in str(exc.value)
